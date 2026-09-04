from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.models import (
    Asset,
    AssetTransaction,
    Bank,
    BankTransaction,
    CreditCard,
    CreditCardTransaction,
    Saving,
    SavingTransaction,
    Transaction,
)
from app.utils.transaction_export_util import (
    CATEGORY_ORDER,
    TRANSACTION_HEADERS,
    TRANSFER_HEADERS,
    get_transaction_export_rows,
    normalize_export_categories,
)


# ============================================================
# CATEGORY NORMALIZATION
# ============================================================

def test_normalize_export_categories_defaults_to_all():
    assert normalize_export_categories() == [
        "banks",
        "savings",
        "assets",
        "credit_cards",
    ]


def test_normalize_export_categories_accepts_all():
    assert normalize_export_categories(
        ["all"]
    ) == list(CATEGORY_ORDER)


@pytest.mark.parametrize(
    "value, expected",
    [
        (
            "banks",
            ["banks"],
        ),
        (
            "savings",
            ["savings"],
        ),
        (
            "assets",
            ["assets"],
        ),
        (
            "credit_cards",
            ["credit_cards"],
        ),
    ],
)
def test_normalize_export_categories_accepts_each_category(
    value,
    expected,
):
    assert normalize_export_categories(
        [value]
    ) == expected


def test_normalize_export_categories_returns_canonical_order():
    assert normalize_export_categories(
        [
            "credit_cards",
            "assets",
            "banks",
        ]
    ) == [
        "banks",
        "assets",
        "credit_cards",
    ]


def test_normalize_export_categories_accepts_comma_separated_values():
    assert normalize_export_categories(
        ["assets,savings"]
    ) == [
        "savings",
        "assets",
    ]


def test_normalize_export_categories_accepts_repeated_values():
    assert normalize_export_categories(
        [
            "credit_cards",
            "assets",
        ]
    ) == [
        "assets",
        "credit_cards",
    ]


def test_normalize_export_categories_ignores_duplicates():
    assert normalize_export_categories(
        [
            "assets",
            "assets",
            "banks",
        ]
    ) == [
        "banks",
        "assets",
    ]


@pytest.mark.parametrize(
    "value, expected",
    [
        (
            "bank",
            ["banks"],
        ),
        (
            "saving",
            ["savings"],
        ),
        (
            "asset",
            ["assets"],
        ),
        (
            "credit-card",
            ["credit_cards"],
        ),
        (
            "select-all",
            list(CATEGORY_ORDER),
        ),
    ],
)
def test_normalize_export_categories_accepts_aliases(
    value,
    expected,
):
    assert normalize_export_categories(
        [value]
    ) == expected


def test_normalize_export_categories_rejects_unknown_category():
    with pytest.raises(
        ValueError,
        match="Invalid transaction category",
    ):
        normalize_export_categories(
            ["crypto"]
        )


def test_normalize_export_categories_rejects_all_with_other_category():
    with pytest.raises(
        ValueError,
        match="'all' cannot be combined",
    ):
        normalize_export_categories(
            [
                "all",
                "banks",
            ]
        )


# ============================================================
# COMMON TRANSACTION DATA LAYER
# ============================================================

def test_export_data_contains_all_supported_categories(
    db_session,
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id
    )

    assert list(export_data.keys()) == [
        "banks",
        "savings",
        "assets",
        "credit_cards",
    ]


def test_export_data_contains_one_transaction_per_category(
    db_session,
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id
    )

    assert len(
        export_data["banks"]
    ) == 1

    assert len(
        export_data["savings"]
    ) == 1

    assert len(
        export_data["assets"]
    ) == 1

    assert len(
        export_data["credit_cards"]
    ) == 1


def test_export_data_normalizes_bank_transaction(
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id,
        ["banks"],
    )

    row = export_data["banks"][0]

    assert row["id"] == (
        export_test_data[
            "bank_transaction"
        ].id
    )

    assert row["amount"] == 1000
    assert row["transaction_type"] == "income"
    assert row["description"] == (
        "BANK_EXPORT_TEST"
    )
    assert row["category"] == "Income"
    assert row["account_type"] == "Bank"
    assert row["account_name"] == (
        "Export Test Bank"
    )
    assert row["balance_after"] == 2000


def test_export_data_normalizes_saving_transaction(
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id,
        ["savings"],
    )

    row = export_data["savings"][0]

    assert row["amount"] == 500
    assert row["transaction_type"] == "deposit"
    assert row["description"] == (
        "SAVING_EXPORT_TEST"
    )
    assert row["account_type"] == "Saving"
    assert row["account_name"] == (
        "Export Test Saving"
    )
    assert row["balance_after"] == 2500


def test_export_data_normalizes_asset_transaction(
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id,
        ["assets"],
    )

    row = export_data["assets"][0]

    assert row["amount"] == 750
    assert row["transaction_type"] == "deposit"
    assert row["description"] == (
        "ASSET_EXPORT_TEST"
    )
    assert row["account_type"] == "Asset"
    assert row["account_name"] == (
        "Export Test Asset"
    )
    assert row["balance_after"] == 3750


def test_export_data_normalizes_credit_card_transaction(
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id,
        ["credit_cards"],
    )

    row = export_data[
        "credit_cards"
    ][0]

    assert row["amount"] == -250
    assert row["transaction_type"] == (
        "expense"
    )
    assert row["description"] == (
        "CARD_EXPORT_TEST"
    )
    assert row["account_type"] == (
        "Credit Card"
    )
    assert row["account_name"] == (
        "Export Test Credit Card"
    )
    assert row["balance_after"] == 1250


def test_export_data_filters_to_requested_categories(
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id,
        [
            "banks",
            "assets",
        ],
    )

    assert list(export_data.keys()) == [
        "banks",
        "assets",
    ]

    assert len(
        export_data["banks"]
    ) == 1

    assert len(
        export_data["assets"]
    ) == 1


def test_export_data_does_not_include_unselected_categories(
    test_user,
    export_test_data,
):
    export_data = get_transaction_export_rows(
        test_user.id,
        ["banks"],
    )

    assert "savings" not in export_data
    assert "assets" not in export_data
    assert "credit_cards" not in export_data


def test_export_data_isolated_to_current_user(
    authenticated_client,
    test_user,
    other_user,
    db_session,
    export_test_data,
):
    other_bank = Bank(
        name="Other Export Bank",
        user_id=other_user.id,
        balance=100,
    )

    db_session.add(other_bank)
    db_session.commit()

    other_transaction = BankTransaction(
        bank_id=other_bank.id,
        user_id=other_user.id,
        amount=9999,
        description="OTHER_USER_EXPORT_TEST",
        category="Other",
        transaction_type="income",
        bank_balance_after=10099,
    )

    db_session.add(
        other_transaction
    )
    db_session.commit()

    export_data = get_transaction_export_rows(
        test_user.id
    )

    exported_descriptions = [
        row["description"]
        for row in export_data["banks"]
    ]

    assert "BANK_EXPORT_TEST" in (
        exported_descriptions
    )

    assert "OTHER_USER_EXPORT_TEST" not in (
        exported_descriptions
    )


def test_export_data_does_not_leak_other_users_account_name(
    test_user,
    other_user,
    db_session,
):
    other_bank = Bank(
        name="Other User Secret Bank",
        user_id=other_user.id,
        balance=100,
    )

    db_session.add(
        other_bank
    )
    db_session.commit()

    compromised_transaction = BankTransaction(
        bank_id=other_bank.id,
        user_id=test_user.id,
        amount=123,
        description="ISOLATION_TEST",
        category="Test",
        transaction_type="income",
        bank_balance_after=123,
    )

    db_session.add(
        compromised_transaction
    )
    db_session.commit()

    export_data = get_transaction_export_rows(
        test_user.id,
        ["banks"],
    )

    row = export_data[
        "banks"
    ][0]

    assert row["description"] == (
        "ISOLATION_TEST"
    )

    assert row["account_name"] == ""


# ============================================================
# EXCEL HELPERS
# ============================================================

def _load_workbook_from_response(
    response,
):
    return load_workbook(
        filename=BytesIO(
            response.data
        ),
        read_only=True,
    )


def _get_sheet_rows(
    workbook,
    sheet_name,
):
    return list(
        workbook[sheet_name].iter_rows(
            values_only=True
        )
    )


# ============================================================
# EXCEL ROUTE
# ============================================================

def test_export_excel_returns_valid_workbook(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
    )

    assert response.status_code == 200

    assert response.mimetype == (
        "application/"
        "vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )

    assert response.data.startswith(
        b"PK"
    )


def test_export_excel_default_contains_all_categories(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
    )

    workbook = _load_workbook_from_response(
        response
    )

    assert workbook.sheetnames == [
        "Banks",
        "Savings",
        "Assets",
        "Credit Cards",
        "Transfers",
    ]


def test_export_excel_all_matches_default(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
        "?categories=all"
    )

    assert response.status_code == 200

    workbook = _load_workbook_from_response(
        response
    )

    assert workbook.sheetnames == [
        "Banks",
        "Savings",
        "Assets",
        "Credit Cards",
        "Transfers",
    ]


@pytest.mark.parametrize(
    "category, sheet_name, marker",
    [
        (
            "banks",
            "Banks",
            "BANK_EXPORT_TEST",
        ),
        (
            "savings",
            "Savings",
            "SAVING_EXPORT_TEST",
        ),
        (
            "assets",
            "Assets",
            "ASSET_EXPORT_TEST",
        ),
        (
            "credit_cards",
            "Credit Cards",
            "CARD_EXPORT_TEST",
        ),
    ],
)
def test_export_excel_single_category(
    authenticated_client,
    test_user,
    export_test_data,
    category,
    sheet_name,
    marker,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
        f"?categories={category}"
    )

    assert response.status_code == 200

    workbook = _load_workbook_from_response(
        response
    )

    assert workbook.sheetnames == [
        sheet_name,
        "Transfers",
    ]

    rows = _get_sheet_rows(
        workbook,
        sheet_name,
    )

    assert rows[0] == tuple(
        TRANSACTION_HEADERS
    )

    assert len(rows) == 2

    row_text = " ".join(
        str(value)
        for value in rows[1]
        if value is not None
    )

    assert marker in row_text


def test_export_excel_multiple_categories(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
        "?categories=banks,assets"
    )

    assert response.status_code == 200

    workbook = _load_workbook_from_response(
        response
    )

    assert workbook.sheetnames == [
        "Banks",
        "Assets",
        "Transfers",
    ]


def test_export_excel_category_data_is_not_mixed(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
        "?categories=assets"
    )

    workbook = _load_workbook_from_response(
        response
    )

    rows = _get_sheet_rows(
        workbook,
        "Assets",
    )

    exported_text = "\n".join(
        " ".join(
            str(value)
            for value in row
            if value is not None
        )
        for row in rows
    )

    assert "ASSET_EXPORT_TEST" in (
        exported_text
    )

    assert "BANK_EXPORT_TEST" not in (
        exported_text
    )

    assert "SAVING_EXPORT_TEST" not in (
        exported_text
    )

    assert "CARD_EXPORT_TEST" not in (
        exported_text
    )


def test_export_excel_transfers_sheet_is_empty(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
    )

    workbook = _load_workbook_from_response(
        response
    )

    rows = _get_sheet_rows(
        workbook,
        "Transfers",
    )

    assert len(rows) == 1

    assert rows[0] == tuple(
        TRANSFER_HEADERS
    )


def test_export_excel_with_no_history(
    authenticated_client,
    test_user,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
    )

    assert response.status_code == 200

    workbook = _load_workbook_from_response(
        response
    )

    for sheet_name in [
        "Banks",
        "Savings",
        "Assets",
        "Credit Cards",
    ]:
        rows = _get_sheet_rows(
            workbook,
            sheet_name,
        )

        assert len(rows) == 1


def test_export_excel_rejects_invalid_category(
    authenticated_client,
    test_user,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
        "?categories=invalid"
    )

    assert response.status_code == 400

    data = response.get_json()

    assert (
        "Invalid transaction category"
        in data["error"]
    )

    assert data[
        "allowed_categories"
    ] == [
        "all",
        "banks",
        "savings",
        "assets",
        "credit_cards",
    ]


def test_export_excel_rejects_all_with_other_category(
    authenticated_client,
    test_user,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/excel"
        "?categories=all,banks"
    )

    assert response.status_code == 400


def test_export_excel_requires_authentication(
    client,
    test_user,
):
    response = client.get(
        f"/users/{test_user.id}/transactions/export/excel"
    )

    assert response.status_code == 401


def test_export_excel_rejects_other_user(
    authenticated_client,
    other_user,
):
    response = authenticated_client.get(
        f"/users/{other_user.id}/transactions/export/excel"
    )

    assert response.status_code == 403

    data = response.get_json()

    assert data["error"] == "Unauthorized"


def test_export_excel_rejects_nonexistent_user(
    authenticated_client,
):
    response = authenticated_client.get(
        "/users/999999/transactions/export/excel"
    )

    assert response.status_code == 404

    data = response.get_json()

    assert data["error"] == "User not found"


# ============================================================
# PDF HELPERS
# ============================================================

def _extract_pdf_text(
    response,
):
    reader = PdfReader(
        BytesIO(response.data)
    )

    return "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )


# ============================================================
# PDF ROUTE
# ============================================================

def test_export_pdf_returns_valid_pdf(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
    )

    assert response.status_code == 200
    assert response.mimetype == (
        "application/pdf"
    )
    assert response.data.startswith(
        b"%PDF"
    )


def test_export_pdf_default_contains_all_sections(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
    )

    text = _extract_pdf_text(
        response
    )

    assert "Banks" in text
    assert "Savings" in text
    assert "Assets" in text
    assert "Credit Cards" in text
    assert "Transfers" in text


@pytest.mark.parametrize(
    "category, expected_section, marker",
    [
        (
            "banks",
            "Banks",
            "BANK_EXPORT_TEST",
        ),
        (
            "savings",
            "Savings",
            "SAVING_EXPORT_TEST",
        ),
        (
            "assets",
            "Assets",
            "ASSET_EXPORT_TEST",
        ),
        (
            "credit_cards",
            "Credit Cards",
            "CARD_EXPORT_TEST",
        ),
    ],
)
def test_export_pdf_single_category(
    authenticated_client,
    test_user,
    export_test_data,
    category,
    expected_section,
    marker,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
        f"?categories={category}"
    )

    assert response.status_code == 200

    text = _extract_pdf_text(
        response
    )

    assert expected_section in text
    assert marker in text


def test_export_pdf_multiple_categories(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
        "?categories=banks,assets"
    )

    assert response.status_code == 200

    text = _extract_pdf_text(
        response
    )

    assert "Banks" in text
    assert "Assets" in text
    assert "BANK_EXPORT_TEST" in text
    assert "ASSET_EXPORT_TEST" in text

    assert "SAVING_EXPORT_TEST" not in (
        text
    )

    assert "CARD_EXPORT_TEST" not in (
        text
    )


def test_export_pdf_category_data_is_not_mixed(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
        "?categories=assets"
    )

    text = _extract_pdf_text(
        response
    )

    assert "ASSET_EXPORT_TEST" in text

    assert "BANK_EXPORT_TEST" not in text
    assert "SAVING_EXPORT_TEST" not in text
    assert "CARD_EXPORT_TEST" not in text


def test_export_pdf_transfers_remain_empty(
    authenticated_client,
    test_user,
    export_test_data,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
    )

    text = _extract_pdf_text(
        response
    )

    assert "Transfers" in text
    assert "No transfer records" in text


def test_export_pdf_with_no_history(
    authenticated_client,
    test_user,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
    )

    assert response.status_code == 200
    assert response.mimetype == (
        "application/pdf"
    )
    assert response.data.startswith(
        b"%PDF"
    )

    text = _extract_pdf_text(
        response
    )

    assert "No transaction records" in text


def test_export_pdf_rejects_invalid_category(
    authenticated_client,
    test_user,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
        "?categories=invalid"
    )

    assert response.status_code == 400

    data = response.get_json()

    assert (
        "Invalid transaction category"
        in data["error"]
    )


def test_export_pdf_rejects_all_with_other_category(
    authenticated_client,
    test_user,
):
    response = authenticated_client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
        "?categories=all,banks"
    )

    assert response.status_code == 400


def test_export_pdf_requires_authentication(
    client,
    test_user,
):
    response = client.get(
        f"/users/{test_user.id}/transactions/export/pdf"
    )

    assert response.status_code == 401


def test_export_pdf_rejects_other_user(
    authenticated_client,
    other_user,
):
    response = authenticated_client.get(
        f"/users/{other_user.id}/transactions/export/pdf"
    )

    assert response.status_code == 403

    data = response.get_json()

    assert data["error"] == "Unauthorized"


def test_export_pdf_rejects_nonexistent_user(
    authenticated_client,
):
    response = authenticated_client.get(
        "/users/999999/transactions/export/pdf"
    )

    assert response.status_code == 404

    data = response.get_json()

    assert data["error"] == "User not found" 