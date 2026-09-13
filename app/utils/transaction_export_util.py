from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .. import db
from ..models import (
    Asset,
    AssetTransaction,
    Bank,
    BankTransaction,
    CreditCard,
    CreditCardTransaction,
    Saving,
    SavingTransaction,
)


# ============================================================
# EXPORT CATEGORY CONFIGURATION
# ============================================================

TRANSACTION_EXPORT_CONFIG = {
    "banks": {
        "label": "Banks",
        "account_type": "Bank",
        "transaction_model": BankTransaction,
        "account_model": Bank,
        "account_id_field": "bank_id",
        "balance_field": "bank_balance_after",
    },
    "savings": {
        "label": "Savings",
        "account_type": "Saving",
        "transaction_model": SavingTransaction,
        "account_model": Saving,
        "account_id_field": "saving_id",
        "balance_field": "saving_balance_after",
    },
    "assets": {
        "label": "Assets",
        "account_type": "Asset",
        "transaction_model": AssetTransaction,
        "account_model": Asset,
        "account_id_field": "asset_id",
        "balance_field": "asset_balance_after",
    },
    "credit_cards": {
        "label": "Credit Cards",
        "account_type": "Credit Card",
        "transaction_model": CreditCardTransaction,
        "account_model": CreditCard,
        "account_id_field": "credit_card_id",
        "balance_field": "card_balance_after",
    },
}

CATEGORY_ORDER = tuple(
    TRANSACTION_EXPORT_CONFIG.keys()
)

CATEGORY_ALIASES = {
    "bank": "banks",
    "saving": "savings",
    "asset": "assets",
    "credit_card": "credit_cards",
    "credit-card": "credit_cards",
    "creditcard": "credit_cards",
    "credit-cards": "credit_cards",
    "select_all": "all",
    "select-all": "all",
}

TRANSACTION_HEADERS = [
    "ID",
    "Date",
    "Amount",
    "Transaction Type",
    "Description",
    "Category",
    "Account Type",
    "Account Name",
    "Balance After",
]

TRANSFER_HEADERS = [
    "ID",
    "Date",
    "Amount",
    "From Account Type",
    "From Account ID",
    "To Account Type",
    "To Account ID",
    "Description",
    "Fee",
]


# ============================================================
# CATEGORY NORMALIZATION
# ============================================================

def normalize_export_categories(values=None):
    """
    Validate and normalize selected transaction categories.

    Supported values:

        all
        banks
        savings
        assets
        credit_cards

    The function accepts both comma-separated values and
    repeated query parameters.
    """

    if not values:
        return list(CATEGORY_ORDER)

    tokens = []

    for value in values:
        tokens.extend(
            token.strip().lower()
            for token in value.split(",")
            if token.strip()
        )

    if not tokens:
        return list(CATEGORY_ORDER)

    normalized = [
        CATEGORY_ALIASES.get(
            token,
            token,
        )
        for token in tokens
    ]

    if "all" in normalized:
        if len(normalized) > 1:
            raise ValueError(
                "'all' cannot be combined with other categories"
            )

        return list(CATEGORY_ORDER)

    unknown = [
        category
        for category in normalized
        if category not in TRANSACTION_EXPORT_CONFIG
    ]

    if unknown:
        allowed = ", ".join(
            [
                "all",
                *CATEGORY_ORDER,
            ]
        )

        raise ValueError(
            f"Invalid transaction category: "
            f"{unknown[0]}. "
            f"Allowed values: {allowed}"
        )

    return [
        category
        for category in CATEGORY_ORDER
        if category in normalized
    ]


# ============================================================
# COMMON TRANSACTION DATA LAYER
# ============================================================

def get_transaction_export_rows(
    user_id,
    categories=None,
):
    """
    Return normalized transaction data grouped by category.

    This layer intentionally knows nothing about Excel or PDF.
    Both export formats consume the same normalized data.
    """

    selected_categories = normalize_export_categories(
        categories
    )

    export_data = {
        category: []
        for category in selected_categories
    }

    for category in selected_categories:
        config = TRANSACTION_EXPORT_CONFIG[
            category
        ]

        transaction_model = config[
            "transaction_model"
        ]

        account_model = config[
            "account_model"
        ]

        account_id_column = getattr(
            transaction_model,
            config["account_id_field"],
        )

        balance_field = config[
            "balance_field"
        ]

        records = (
            db.session
            .query(
                transaction_model,
                account_model.name,
            )
            .outerjoin(
                account_model,
                (
                    account_id_column
                    == account_model.id
                )
                & (
                    account_model.user_id
                    == user_id
                ),
            )
            .filter(
                transaction_model.user_id
                == user_id
            )
            .order_by(
                transaction_model.date.asc(),
                transaction_model.id.asc(),
            )
            .all()
        )

        export_data[category] = [
            {
                "id": transaction.id,
                "date": transaction.date,
                "amount": transaction.amount,
                "transaction_type": (
                    transaction.transaction_type
                    or ""
                ),
                "description": (
                    transaction.description
                    or ""
                ),
                "category": (
                    transaction.category
                    or ""
                ),
                "account_type": (
                    config["account_type"]
                ),
                "account_name": (
                    account_name or ""
                ),
                "balance_after": getattr(
                    transaction,
                    balance_field,
                ),
            }
            for transaction, account_name
            in records
        ]

    return export_data


# ============================================================
# EXCEL
# ============================================================

def _style_excel_sheet(sheet):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E1F2",
    )

    header_font = Font(
        bold=True
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in sheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    header_columns = {
        cell.value: cell.column_letter
        for cell in sheet[1]
    }

    date_column = header_columns.get(
        "Date"
    )

    amount_column = header_columns.get(
        "Amount"
    )

    balance_column = header_columns.get(
        "Balance After"
    )

    if date_column:
        for cell in sheet[date_column][1:]:
            if cell.value is not None:
                cell.number_format = (
                    "yyyy-mm-dd hh:mm:ss"
                )

    if amount_column:
        for cell in sheet[amount_column][1:]:
            if cell.value is not None:
                cell.number_format = "0.00"

    if balance_column:
        for cell in sheet[balance_column][1:]:
            if cell.value is not None:
                cell.number_format = "0.00"

    for column_cells in sheet.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column_cells
        )

        column_letter = (
            column_cells[0].column_letter
        )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            40,
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _append_transaction_sheet(
    workbook,
    title,
    rows,
):
    sheet = workbook.create_sheet(title)

    sheet.append(
        TRANSACTION_HEADERS
    )

    for row in rows:
        excel_date = row["date"]

        if (
            excel_date is not None
            and excel_date.tzinfo is not None
        ):
            excel_date = excel_date.replace(
                tzinfo=None
            )

        sheet.append([
            row["id"],
            excel_date,
            row["amount"],
            row["transaction_type"],
            row["description"],
            row["category"],
            row["account_type"],
            row["account_name"],
            row["balance_after"],
        ])

    _style_excel_sheet(sheet)

    return sheet


def _append_empty_transfer_sheet(
    workbook,
):
    sheet = workbook.create_sheet(
        "Transfers"
    )

    sheet.append(
        TRANSFER_HEADERS
    )

    _style_excel_sheet(sheet)

    return sheet


def create_transaction_excel(
    user_id,
    categories=None,
):
    selected_categories = (
        normalize_export_categories(
            categories
        )
    )

    export_data = (
        get_transaction_export_rows(
            user_id,
            selected_categories,
        )
    )

    workbook = Workbook()

    workbook.remove(
        workbook.active
    )

    for category in selected_categories:
        config = (
            TRANSACTION_EXPORT_CONFIG[
                category
            ]
        )

        _append_transaction_sheet(
            workbook,
            config["label"],
            export_data[category],
        )

    # Transfers deliberately remain empty.
    _append_empty_transfer_sheet(
        workbook
    )

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


# ============================================================
# PDF
# ============================================================

def _pdf_cell(
    value,
    style,
):
    text = (
        ""
        if value is None
        else str(value)
    )

    return Paragraph(
        escape(text),
        style,
    )


def _build_pdf_transaction_table(
    rows,
    styles,
):
    data = [[
        _pdf_cell(
            header,
            styles["TableHeader"],
        )
        for header
        in TRANSACTION_HEADERS
    ]]

    for row in rows:
        date_value = (
            row["date"].strftime(
                "%Y-%m-%d %H:%M:%S"
            )
            if row["date"]
            else ""
        )

        data.append([
            _pdf_cell(
                row["id"],
                styles["TableCell"],
            ),
            _pdf_cell(
                date_value,
                styles["TableCell"],
            ),
            _pdf_cell(
                row["amount"],
                styles["TableCell"],
            ),
            _pdf_cell(
                row["transaction_type"],
                styles["TableCell"],
            ),
            _pdf_cell(
                row["description"],
                styles["TableCell"],
            ),
            _pdf_cell(
                row["category"],
                styles["TableCell"],
            ),
            _pdf_cell(
                row["account_type"],
                styles["TableCell"],
            ),
            _pdf_cell(
                row["account_name"],
                styles["TableCell"],
            ),
            _pdf_cell(
                row["balance_after"],
                styles["TableCell"],
            ),
        ])

    if len(data) == 1:
        data.append([
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "No transaction records",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
            _pdf_cell(
                "",
                styles["TableCell"],
            ),
        ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            11 * mm,
            28 * mm,
            18 * mm,
            22 * mm,
            55 * mm,
            25 * mm,
            23 * mm,
            35 * mm,
            25 * mm,
        ],
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey,
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey,
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP",
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
        ])
    )

    return table


def _build_empty_transfer_table(
    styles,
):
    data = [[
        Paragraph(
            header,
            styles["TableHeader"],
        )
        for header
        in TRANSFER_HEADERS
    ], [
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
        Paragraph(
            "No transfer records",
            styles["TableCell"],
        ),
        Paragraph(
            "",
            styles["TableCell"],
        ),
    ]]

    table = Table(
        data,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey,
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey,
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP",
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
        ])
    )

    return table


def create_transaction_pdf(
    user_id,
    categories=None,
):
    selected_categories = (
        normalize_export_categories(
            categories
        )
    )

    export_data = (
        get_transaction_export_rows(
            user_id,
            selected_categories,
        )
    )

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=8 * mm,
        leftMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="TableHeader",
            parent=styles["BodyText"],
            fontSize=6.5,
            leading=7.5,
            fontName="Helvetica-Bold",
        )
    )

    styles.add(
        ParagraphStyle(
            name="TableCell",
            parent=styles["BodyText"],
            fontSize=6.5,
            leading=7.5,
        )
    )

    elements = [
        Paragraph(
            "PPA Transaction History Backup",
            styles["Title"],
        ),
        Spacer(
            1,
            4 * mm,
        ),
    ]

    for index, category in enumerate(
        selected_categories
    ):
        config = (
            TRANSACTION_EXPORT_CONFIG[
                category
            ]
        )

        elements.append(
            Paragraph(
                config["label"],
                styles["Heading2"],
            )
        )

        elements.append(
            _build_pdf_transaction_table(
                export_data[category],
                styles,
            )
        )

        if index < (
            len(selected_categories) - 1
        ):
            elements.append(
                Spacer(
                    1,
                    6 * mm,
                )
            )

    # Transfers deliberately remain empty.
    elements.append(
        Spacer(
            1,
            6 * mm,
        )
    )

    elements.append(
        Paragraph(
            "Transfers",
            styles["Heading2"],
        )
    )

    elements.append(
        _build_empty_transfer_table(
            styles
        )
    )

    document.build(elements)

    output.seek(0)

    return output